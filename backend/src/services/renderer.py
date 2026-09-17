"""Rendering service.

Strategy: Always build a clean, professional, JD-tailored resume from scratch.
The resume object already contains the user's data + applied patches from the
generator (summary rewritten, skills reordered, bullets upgraded). This
service just renders it cleanly as a fresh DOCX or PDF.

Template-preserving "patch the original" strategies were removed because
they caused content repetition and unreliable matching. A fresh build is
more reliable and gives a consistent, ATS-friendly output every time.
"""
import os
import subprocess
from datetime import datetime, timezone
from pathlib import Path
from typing import Optional

from ..models.resume import ResumeSchema
from ..services.latex_generator import generate_latex
from ..services.jakes_template import _looks_like_award

OUTPUT_DIR = Path(os.getenv("OUTPUT_DIR", "outputs"))
OUTPUT_DIR.mkdir(exist_ok=True)


# ─────────────────────────────────────────────────────────────────────────────
# Public API
# ─────────────────────────────────────────────────────────────────────────────

def render_docx(
    resume: ResumeSchema,
    original_path: Optional[str] = None,
    original_suffix: Optional[str] = None,
    original_bullets: Optional[dict] = None,
    original_summary: Optional[str] = None,
) -> str:
    """Build a clean, professional DOCX from the tailored resume data.

    The original_*, original_suffix parameters are accepted for backward
    compatibility but ignored: this renderer always drafts a fresh resume
    from the structured data, never patches the uploaded file.
    """
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    out_path  = OUTPUT_DIR / f"resume_{timestamp}.docx"
    return _render_docx_scratch(resume, str(out_path))


def render_pdf(
    resume: ResumeSchema,
    original_path: Optional[str] = None,
    original_suffix: Optional[str] = None,
    original_bullets: Optional[dict] = None,
    original_summary: Optional[str] = None,
) -> str:
    """Build a clean, professional PDF from the tailored resume data.

    The original_* parameters are accepted for backward compatibility but
    ignored: this renderer always drafts a fresh resume from the
    structured data, never patches the uploaded file.
    """
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    out_path  = OUTPUT_DIR / f"resume_{timestamp}.pdf"
    return _render_pdf_scratch(resume, str(out_path))


# ─────────────────────────────────────────────────────────────────────────────
# DOCX scratch renderer
# ─────────────────────────────────────────────────────────────────────────────

def _render_docx_scratch(resume: ResumeSchema, out_path: str) -> str:
    """Build a clean professional DOCX from the resume data structure."""
    try:
        from docx import Document
        from docx.shared import Inches, Pt, RGBColor
        from docx.oxml.ns import qn
        from docx.oxml import OxmlElement

        doc = Document()
        for section in doc.sections:
            section.top_margin    = Inches(0.60)
            section.bottom_margin = Inches(0.60)
            section.left_margin   = Inches(0.70)
            section.right_margin  = Inches(0.70)

        NAVY  = RGBColor(30, 58, 95)
        DARK  = RGBColor(25, 25, 25)
        MUTED = RGBColor(100, 100, 100)

        style = doc.styles["Normal"]
        style.paragraph_format.space_before = Pt(0)
        style.paragraph_format.space_after  = Pt(0)

        def add_rule(paragraph):
            pPr  = paragraph._p.get_or_add_pPr()
            pBdr = OxmlElement("w:pBdr")
            bot  = OxmlElement("w:bottom")
            bot.set(qn("w:val"), "single"); bot.set(qn("w:sz"), "6")
            bot.set(qn("w:space"), "1");    bot.set(qn("w:color"), "1E3A5F")
            pBdr.append(bot); pPr.append(pBdr)

        def heading(title: str):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(9)
            p.paragraph_format.space_after  = Pt(2)
            r = p.add_run(title.upper())
            r.font.name = "Calibri"; r.font.size = Pt(10.5)
            r.font.bold = True;      r.font.color.rgb = NAVY
            add_rule(p)

        def run(para, text, size=9.5, bold=False, color=None):
            r = para.add_run(text)
            r.font.name = "Calibri"; r.font.size = Pt(size)
            r.font.bold = bold;      r.font.color.rgb = color or DARK

        # Name
        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(1)
        run(p, resume.contact.name or "Resume", 20, True, NAVY)

        # Contact
        parts = [x for x in [resume.contact.email, resume.contact.phone] + (resume.contact.links or []) if x]
        if parts:
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(6)
            run(p, "  |  ".join(parts), 9.0, color=MUTED)

        # Summary (already JD-tailored by the generator)
        if resume.summary:
            heading("Professional Summary")
            p = doc.add_paragraph()
            p.paragraph_format.space_after = Pt(3)
            run(p, resume.summary.strip())

        # Skills (already reordered by the generator for ATS)
        if resume.skills:
            heading("Technical Skills")
            cat_map: dict[str, list[str]] = {}
            for s in resume.skills:
                cat_map.setdefault((s.category or "General").title(), []).append(s.name)
            for cat, names in cat_map.items():
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                run(p, f"{cat}: ", bold=True)
                run(p, ", ".join(names[:12]))

        # Experience (bullets already upgraded by the generator)
        if resume.experience:
            heading("Work Experience")
            for exp in resume.experience:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(5)
                p.paragraph_format.space_after  = Pt(1)
                run(p, exp.title or "", 10.5, True)
                if exp.company: run(p, f"  —  {exp.company}", 10.5)
                d = _fmt_dates(exp.start_date, exp.end_date)
                if d: run(p, f"   ({d})", 9.0, color=MUTED)
                for b in exp.bullets:
                    p2 = doc.add_paragraph(style="List Bullet")
                    p2.paragraph_format.space_after = Pt(1)
                    run(p2, b.text)

        # Education
        if resume.education:
            heading("Education")
            for edu in resume.education:
                p = doc.add_paragraph()
                p.paragraph_format.space_after = Pt(2)
                line = edu.school
                if edu.degree: line += f"  —  {edu.degree}"
                if edu.dates:  line += f" ({edu.dates})"
                run(p, line, 10.0)

        # Projects
        if resume.projects:
            heading("Projects")
            for proj in resume.projects:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(3)
                p.paragraph_format.space_after  = Pt(1)
                run(p, proj.name, 10.0, True)
                if proj.description:
                    p2 = doc.add_paragraph()
                    p2.paragraph_format.space_after = Pt(1)
                    run(p2, proj.description)
                if proj.technologies:
                    p2 = doc.add_paragraph()
                    p2.paragraph_format.space_after = Pt(2)
                    run(p2, "Tech: " + ", ".join(proj.technologies), 9.0, color=MUTED)

        # Certifications & Achievements
        certs = [x for x in resume.certifications if not _looks_like_award(x.name, x.issuer)]
        achievements = [x for x in resume.certifications if _looks_like_award(x.name, x.issuer)]

        if certs:
            heading("Certifications")
            for cert in certs:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                ct = cert.name
                if cert.issuer: ct += f"  —  {cert.issuer}"
                if cert.date:   ct += f" ({cert.date})"
                run(p, ct)

        if achievements:
            heading("Achievements")
            for ach in achievements:
                p = doc.add_paragraph(style="List Bullet")
                p.paragraph_format.space_after = Pt(1)
                ct = ach.name
                if ach.issuer: ct += f"  —  {ach.issuer}"
                if ach.date:   ct += f" ({ach.date})"
                run(p, ct)

        doc.save(out_path)
        return out_path

    except Exception as e:
        # Last resort: write HTML so the user still gets something readable
        print(f"[Renderer] Scratch DOCX failed: {e}")
        Path(out_path).write_text(_build_html(resume), encoding="utf-8")
        return out_path


# ─────────────────────────────────────────────────────────────────────────────
# PDF scratch renderer
# ─────────────────────────────────────────────────────────────────────────────

def _render_pdf_scratch(resume: ResumeSchema, out_path: str) -> str:
    """Build a clean ATS-safe PDF from the resume data structure.

    Tries PyMuPDF first, then WeasyPrint, then falls back to HTML.
    """
    try:
        try:
            import pymupdf as fitz  # type: ignore
        except ImportError:
            import fitz  # type: ignore
        return _build_pdf_pymupdf(resume, out_path)
    except Exception as e:
        print(f"[Renderer] PyMuPDF scratch failed ({e}), trying WeasyPrint")

    try:
        from weasyprint import HTML  # type: ignore
        HTML(string=_build_html(resume)).write_pdf(target=out_path)
        return out_path
    except Exception as e:
        print(f"[Renderer] WeasyPrint failed ({e}), saving HTML")

    html_path = out_path.replace(".pdf", ".html")
    Path(html_path).write_text(_build_html(resume), encoding="utf-8")
    return html_path


def _build_pdf_pymupdf(resume: ResumeSchema, output_path: str) -> str:
    """Render a clean PDF with PyMuPDF using the resume data structure."""
    try:
        import pymupdf as fitz  # type: ignore
    except ImportError:
        import fitz  # type: ignore

    PW, PH = 612.0, 792.0
    ML, MR, MT, MB = 48.0, 48.0, 42.0, 50.0
    CW = PW - ML - MR
    NAVY  = (0.12, 0.23, 0.37)
    DARK  = (0.10, 0.10, 0.10)
    MUTED = (0.43, 0.43, 0.43)

    doc  = fitz.open()
    page = doc.new_page(width=PW, height=PH)
    y    = MT

    def new_pg():
        nonlocal page, y
        page = doc.new_page(width=PW, height=PH); y = MT

    def need(h):
        nonlocal page, y
        if y + h > PH - MB: new_pg()

    def para(txt, fs=9.5, fn="helv", color=DARK, indent=0.0):
        nonlocal y
        if not txt: return
        x = ML + indent; mw = CW - indent
        words = txt.split(); line = ""
        for w in words:
            test = f"{line} {w}".strip()
            if fitz.get_text_length(test, fontname=fn, fontsize=fs) <= mw:
                line = test
            else:
                if line:
                    need(fs+3.5); page.insert_text((x,y+fs),line,fontsize=fs,fontname=fn,color=color); y+=fs+3.5
                line = w
        if line:
            need(fs+3.5); page.insert_text((x,y+fs),line,fontsize=fs,fontname=fn,color=color); y+=fs+3.5

    def sec(title):
        nonlocal y
        need(28); y+=7
        page.insert_text((ML,y+10),title.upper(),fontsize=10,fontname="hebo",color=NAVY); y+=14
        page.draw_line((ML,y),(PW-MR,y),color=NAVY,width=0.8); y+=5

    # Header
    need(30)
    page.insert_text((ML,y+18),resume.contact.name or "Candidate",fontsize=19,fontname="hebo",color=NAVY); y+=23
    parts=[x for x in [resume.contact.email,resume.contact.phone]+(resume.contact.links or []) if x]
    if parts: para("  |  ".join(parts),9.0,"helv",MUTED); y+=2

    # Summary (JD-tailored)
    if resume.summary:
        sec("Professional Summary"); para(resume.summary.strip()); y+=2

    # Skills (reordered for ATS)
    if resume.skills:
        sec("Technical Skills")
        cat_map: dict[str,list[str]] = {}
        for s in resume.skills:
            cat_map.setdefault((s.category or "General").title(),[]).append(s.name)
        for cat,names in cat_map.items():
            para(f"{cat}: "+", ".join(names[:12])); y+=1

    # Experience (upgraded bullets)
    if resume.experience:
        sec("Work Experience")
        for exp in resume.experience:
            need(20)
            hdr = exp.title or ""
            if exp.company: hdr += f"  —  {exp.company}"
            d = _fmt_dates(exp.start_date, exp.end_date)
            page.insert_text((ML,y+10),hdr,fontsize=10,fontname="hebo",color=DARK)
            if d:
                dw=fitz.get_text_length(d,fontname="helv",fontsize=8.5)
                page.insert_text((PW-MR-dw,y+9.5),d,fontsize=8.5,fontname="helv",color=MUTED)
            y+=13
            for b in exp.bullets:
                para(f"•  {b.text}",9.0,"helv",DARK,7.0)
            y+=3

    # Education
    if resume.education:
        sec("Education")
        for edu in resume.education:
            line=edu.school
            if edu.degree: line+=f"  —  {edu.degree}"
            if edu.dates:  line+=f" ({edu.dates})"
            para(line)

    # Projects
    if resume.projects:
        sec("Projects")
        for proj in resume.projects:
            need(18)
            para(proj.name,9.5,"hebo",DARK)
            if proj.description: para(proj.description,9.0,"helv",DARK,5.0)
            if proj.technologies: para("Tech: "+", ".join(proj.technologies),8.5,"helv",MUTED,5.0)
            y+=2

    # Certifications & Achievements
    certs = [x for x in resume.certifications if not _looks_like_award(x.name, x.issuer)]
    achievements = [x for x in resume.certifications if _looks_like_award(x.name, x.issuer)]

    if certs:
        sec("Certifications")
        for cert in certs:
            line=f"•  {cert.name}"
            if cert.issuer: line+=f"  —  {cert.issuer}"
            if cert.date:   line+=f" ({cert.date})"
            para(line,9.0,"helv",DARK)

    if achievements:
        sec("Achievements")
        for ach in achievements:
            line=f"•  {ach.name}"
            if ach.issuer: line+=f"  —  {ach.issuer}"
            if ach.date:   line+=f" ({ach.date})"
            para(line,9.0,"helv",DARK)

    doc.save(output_path); doc.close()
    return output_path


# ─────────────────────────────────────────────────────────────────────────────
# HTML fallback (used when both DOCX and PDF backends fail)
# ─────────────────────────────────────────────────────────────────────────────

def _build_html(resume: ResumeSchema) -> str:
    def esc(s):
        return (s or "").replace("&","&amp;").replace("<","&lt;").replace(">","&gt;")

    cparts=[esc(resume.contact.email),esc(resume.contact.phone)]+[f'<a href="https://{esc(l)}">{esc(l)}</a>' for l in (resume.contact.links or [])]
    contact_html="  |  ".join(p for p in cparts if p)

    skill_html=""
    if resume.skills:
        cm: dict[str,list[str]]={}
        for s in resume.skills: cm.setdefault((s.category or "General").title(),[]).append(s.name)
        skill_html="<br>".join(f"<b>{esc(c)}:</b> {esc(', '.join(ns[:12]))}" for c,ns in cm.items())

    exp_html=""
    for exp in resume.experience:
        d=_fmt_dates(exp.start_date,exp.end_date)
        hdr=f"<b>{esc(exp.title)}</b>"+( f" &mdash; <i>{esc(exp.company)}</i>" if exp.company else "")
        ds=f'<span style="color:#666;font-size:9pt">{esc(d)}</span>' if d else ""
        buls="".join(f"<li>{esc(b.text)}</li>" for b in exp.bullets)
        exp_html+=f'<div style="margin-bottom:8px"><div style="display:flex;justify-content:space-between">{hdr}{ds}</div><ul style="margin:3px 0 0 16px;padding:0">{buls}</ul></div>'

    edu_html="".join(f"<p style='margin:2px 0'>{esc(e.school)}{(' &mdash; '+esc(e.degree)) if e.degree else ''}{(' ('+esc(e.dates)+')') if e.dates else ''}</p>" for e in resume.education)

    proj_html=""
    for p in resume.projects:
        proj_html+=f"<p style='margin:4px 0'><b>{esc(p.name)}</b>"
        if p.description: proj_html+=f"<br>{esc(p.description)}"
        if p.technologies: proj_html+=f"<br><i style='color:#555;font-size:9pt'>Tech: {esc(', '.join(p.technologies))}</i>"
        proj_html+="</p>"

    certs = [x for x in resume.certifications if not _looks_like_award(x.name, x.issuer)]
    achievements = [x for x in resume.certifications if _looks_like_award(x.name, x.issuer)]
    cert_html="".join(f"<li>{esc(c.name)}{(' &mdash; '+esc(c.issuer)) if c.issuer else ''}{(' ('+esc(c.date)+')') if c.date else ''}</li>" for c in certs)
    ach_html="".join(f"<li>{esc(a.name)}{(' &mdash; '+esc(a.issuer)) if a.issuer else ''}{(' ('+esc(a.date)+')') if a.date else ''}</li>" for a in achievements)

    def sec(title,body):
        if not body.strip(): return ""
        return f'<div style="margin-top:10px"><h2 style="font-size:10.5pt;color:#1e3a5f;border-bottom:1.2px solid #1e3a5f;padding-bottom:2px;margin-bottom:5px;text-transform:uppercase;letter-spacing:.05em">{esc(title)}</h2>{body}</div>'

    return f"""<!DOCTYPE html><html lang="en"><head><meta charset="utf-8"><title>{esc(resume.contact.name)} — Resume</title>
<style>*{{box-sizing:border-box;margin:0;padding:0}}body{{font-family:'Segoe UI',Calibri,Arial,sans-serif;margin:.65in .75in;color:#1a1a1a;font-size:10pt;line-height:1.45}}h1{{font-size:20pt;color:#1e3a5f;margin-bottom:2px}}ul{{padding-left:16px;margin-top:3px}}li{{margin-bottom:2px}}a{{color:#1e3a5f;text-decoration:none}}</style></head><body>
<h1>{esc(resume.contact.name or "Resume")}</h1>
<p style="color:#555;font-size:9pt;margin-bottom:10px">{contact_html}</p>
{sec("Professional Summary",f"<p>{esc(resume.summary)}</p>") if resume.summary else ""}
{sec("Technical Skills",f"<p style='font-size:9.5pt'>{skill_html}</p>") if skill_html else ""}
{sec("Work Experience",exp_html) if exp_html else ""}
{sec("Education",edu_html) if edu_html else ""}
{sec("Projects",proj_html) if proj_html else ""}
{sec("Certifications",f"<ul>{cert_html}</ul>") if cert_html else ""}
{sec("Achievements",f"<ul>{ach_html}</ul>") if ach_html else ""}
</body></html>"""


# ─────────────────────────────────────────────────────────────────────────────
# Shared helpers
# ─────────────────────────────────────────────────────────────────────────────

def _fmt_dates(start, end) -> str:
    if not start and not end:
        return ""
    s = start.strftime("%b %Y") if hasattr(start,"strftime") else str(start or "")
    e = "Present" if end is None else (end.strftime("%b %Y") if hasattr(end,"strftime") else str(end))
    if s and e: return f"{s} – {e}"
    return s or e


# ─────────────────────────────────────────────────────────────────────────────
# LaTeX renderer
# ─────────────────────────────────────────────────────────────────────────────

def render_latex(
    resume: ResumeSchema,
    original_path: Optional[str] = None,
    original_suffix: Optional[str] = None,
    original_bullets: Optional[dict] = None,
    original_summary: Optional[str] = None,
) -> tuple[str, str]:
    """Generate LaTeX from the resume data and try to compile to PDF.

    Returns a tuple of (output_path, format) where format is "latex" or "pdf".
    If a LaTeX engine (xelatex/pdflatex) is available, the PDF is compiled
    and the .pdf path is returned.  Otherwise, the .tex file is saved and
    the .tex path is returned for the user to compile locally.
    """
    timestamp = datetime.now(timezone.utc).strftime("%Y%m%d%H%M%S")
    tex_path = str(OUTPUT_DIR / f"resume_{timestamp}.tex")
    pdf_path = str(OUTPUT_DIR / f"resume_{timestamp}.pdf")

    # Generate LaTeX source from grounded resume data
    latex_source = generate_latex(resume)
    Path(tex_path).write_text(latex_source, encoding="utf-8")

    # Try to compile with xelatex or pdflatex
    compiled = False
    for engine in ["xelatex", "pdflatex"]:
        try:
            result = subprocess.run(
                [engine, "-interaction=nonstopmode", "-output-directory",
                 str(OUTPUT_DIR), tex_path],
                capture_output=True, text=True, timeout=30
            )
            # Check if PDF was produced
            if result.returncode == 0 and Path(pdf_path).exists():
                compiled = True
                return pdf_path, "pdf"
        except (FileNotFoundError, subprocess.TimeoutExpired):
            continue

    # If no LaTeX engine succeeded, return the .tex path
    if not compiled:
        return tex_path, "latex"
